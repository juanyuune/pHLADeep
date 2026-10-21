import os, sys, gc, random, pickle, warnings, subprocess, tempfile, math
import numpy as np
import pandas as pd
import openpyxl
import torch
import torch.nn as nn
import torch.nn.functional as F
from sklearn.metrics import roc_curve, auc as sk_auc, confusion_matrix, matthews_corrcoef
from collections import Counter
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.ticker as ticker
import matplotlib.font_manager as fm
warnings.filterwarnings("ignore")

os.environ["TF_FORCE_GPU_ALLOW_GROWTH"] = "true"
os.environ["TF_CPP_MIN_LOG_LEVEL"] = "3"

import tensorflow as tf
gpus = tf.config.list_physical_devices("GPU")
for gpu in gpus:
    try:
        tf.config.experimental.set_memory_growth(gpu, True)
    except RuntimeError:
        pass

import ankh

WORKSPACE    = "/srv/jupyterlab/workspace/HLA1"
CODE_DIR     = os.path.join(WORKSPACE, "code")
RESULTS_DIR  = os.path.join(WORKSPACE, "results", "casestudy_ms")
SUPP_XLSX    = os.path.join(WORKSPACE, "case_study", "41587_2023_1945_MOESM5_ESM.xlsx")
HLA_PSEQ_DIR = os.path.join(WORKSPACE, "hla", "hla_pseudoseq")
HLA_EMB_DIR  = os.path.join(WORKSPACE, "hla", "hla_embeddings_ankh")
NETMHCPAN    = os.path.join(WORKSPACE, "tools", "netMHCpan-4.2", "netMHCpan")

os.makedirs(RESULTS_DIR, exist_ok=True)
sys.path.insert(0, CODE_DIR)

MY_ALLELES = {
    "A*01:01","A*02:01","A*03:01","A*11:01","A*24:02",
    "B*07:02","B*08:01","B*35:01","B*40:01","B*51:01","B*57:01",
}
DECOY_RATIO = 10
RANDOM_SEED = 42
AAs         = list("ACDEFGHIKLMNPQRSTVWY")

ANKH_DIM   = 1536
PEP_LEN    = 9
HLA_LEN    = 34
CONCAT_LEN = PEP_LEN + HLA_LEN
NUM_FILTER = 1024
NUM_HIDDEN = 500
WIN_SIZES  = [2, 3, 5, 9]
NUM_CLASS  = 2

DEVICE = torch.device("cuda" if torch.cuda.is_available() else "cpu")
print(f"Using device: {DEVICE}")

_ws = str(WIN_SIZES)
WEIGHTS_A = os.path.join(WORKSPACE, "results", "stratA",
    f"model_ankh_{ANKH_DIM}dim_{NUM_FILTER}f_{NUM_HIDDEN}h_{_ws}.weights.h5")
WEIGHTS_B = os.path.join(WORKSPACE, "results", "stratB_concat",
    f"model_ankh_{ANKH_DIM}dim_{NUM_FILTER}f_{NUM_HIDDEN}h_{_ws}.pt")
WEIGHTS_C = os.path.join(WORKSPACE, "results", "stratC",
    f"model_ankh_{ANKH_DIM}dim_{NUM_FILTER}f_{NUM_HIDDEN}h_{_ws}.pt")

COLORS = {
    "pHLADeep-Pep":    "#4878CF",
    "pHLADeep-Concat": "#D65F5F",
    "pHLADeep-Attn":   "#1baf7a",
    "NetMHCpan 4.2":   "#E8932A",
    "MHCflurry 2.0":   "#956CB4",
}
LSTYLES = {
    "pHLADeep-Pep":    "-",
    "pHLADeep-Concat": "-",
    "pHLADeep-Attn":   "-",
    "NetMHCpan 4.2":   "--",
    "MHCflurry 2.0":   "--",
}

random.seed(RANDOM_SEED)
np.random.seed(RANDOM_SEED)

print("Loading Gurung 2023 cancer neoepitopes")

wb   = openpyxl.load_workbook(SUPP_XLSX, read_only=True)
seen = set()
positives = []

def _add(pep, allele, length, neoag, source):
    if not (pep and allele):
        return
    pep    = str(pep).strip()
    allele = str(allele).strip()
    key = (pep, allele)
    if key in seen or allele not in MY_ALLELES:
        return
    seen.add(key)
    positives.append({
        "peptide":    pep,
        "allele":     f"HLA-{allele}",
        "length":     int(length) if length else len(pep),
        "neoantigen": neoag,
        "label":      1,
        "source":     source,
    })

ws_tab = wb["Targeted_Data"]
rows   = list(ws_tab.iter_rows(values_only=True))
h      = rows[0]
pi, ai, li, ni = (h.index("Peptide"), h.index("Allele"),
                  h.index("peptide_length"), h.index("Neoantigen"))
for r in rows[1:]:
    _add(r[pi], r[ai], r[li], r[ni], "targeted_MS")

ws_tab2 = wb["Untargeted_Data_Neoantigens"]
rows2   = list(ws_tab2.iter_rows(values_only=True))
h2      = rows2[0]
p2, a2, l2, n2 = (h2.index("Peptide"), h2.index("Allele"),
                   h2.index("peptide_length"), h2.index("Neoantigen"))
for r in rows2[1:]:
    _add(r[p2], r[a2], r[l2], r[n2], "untargeted_MS")

print(f"  Positives: {len(positives)} unique neoepitope-HLA pairs")
print(f"  By allele: {dict(Counter(p['allele'] for p in positives))}")
print(f"  By length: {dict(Counter(p['length'] for p in positives))}")

TRAINED_ALLELES = {"HLA-A*02:01", "HLA-A*03:01"}
positives = [p for p in positives if p["allele"] in TRAINED_ALLELES]
print(f"  Fair subset: {len(positives)} positives ({sorted(TRAINED_ALLELES)})")

print("\nGenerating random decoy negatives")

pos_seqs   = {p["peptide"] for p in positives}
allele_pos = {}
for p in positives:
    allele_pos.setdefault(p["allele"], []).append(p)

negatives = []
for allele, pos_list in allele_pos.items():
    n_needed    = len(pos_list) * DECOY_RATIO
    length_dist = Counter(p["length"] for p in pos_list)
    added = attempts = 0
    while added < n_needed and attempts < n_needed * 300:
        attempts += 1
        L   = random.choices(list(length_dist.keys()), weights=list(length_dist.values()))[0]
        pep = "".join(random.choices(AAs, k=L))
        if pep not in pos_seqs:
            negatives.append({"peptide": pep, "allele": allele,
                               "length": L, "label": 0, "source": "random_decoy"})
            added += 1

print(f"  Negatives: {len(negatives)}")

df = pd.DataFrame(positives + negatives)
df = df.sample(frac=1, random_state=RANDOM_SEED).reset_index(drop=True)
df.to_csv(os.path.join(RESULTS_DIR, "casestudy_dataset.csv"), index=False)
print(f"  Total: {len(df)} ({df['label'].sum()} pos / {(df['label']==0).sum()} neg)")

y_true = df["label"].values

_ankh_mdl = _ankh_tok = None

def get_ankh():
    global _ankh_mdl, _ankh_tok
    if _ankh_mdl is None:
        print("  Loading Ankh-Large ...")
        _ankh_mdl, _ankh_tok = ankh.load_large_model()
        _ankh_mdl = _ankh_mdl.to(DEVICE).eval()
    return _ankh_mdl, _ankh_tok

def embed_one_ankh(seq, target_len):
    model, tok = get_ankh()
    enc = tok.batch_encode_plus(
        [list(seq)], add_special_tokens=True,
        padding=True, is_split_into_words=True, return_tensors="pt")
    with torch.no_grad():
        h = model(input_ids=enc["input_ids"].to(DEVICE),
                  attention_mask=enc["attention_mask"].to(DEVICE),
                  output_hidden_states=True).last_hidden_state
    emb = h[0, 1:len(seq)+1].cpu().numpy().astype(np.float32)
    if emb.shape[0] < target_len:
        pad = np.zeros((target_len - emb.shape[0], ANKH_DIM), dtype=np.float32)
        emb = np.vstack([emb, pad])
    else:
        emb = emb[:target_len]
    return emb

def embed_peptides_ankh(seqs):
    out = []
    for i, s in enumerate(seqs):
        out.append(embed_one_ankh(s, PEP_LEN))
        if (i + 1) % 50 == 0:
            print(f"    {i+1}/{len(seqs)} peptides embedded ...")
            gc.collect()
            if DEVICE.type == "cuda":
                torch.cuda.empty_cache()
    return np.array(out, dtype=np.float32)

def load_hla_ankh_emb(allele_str):
    tag = allele_str.replace("HLA-", "").replace("*", "").replace(":", "")
    candidates = [
        os.path.join(HLA_EMB_DIR, f"{tag}.pkl"),
        os.path.join(HLA_EMB_DIR, f"HLA-{tag}.pkl"),
        os.path.join(HLA_EMB_DIR, f"{tag}.ankh"),
        os.path.join(HLA_EMB_DIR, f"HLA-{tag}.ankh"),
    ]
    for p in candidates:
        if os.path.exists(p):
            with open(p, "rb") as f:
                arr = pickle.load(f)
            arr = np.array(arr, dtype=np.float32)
            if arr.ndim == 3 and arr.shape[0] == 1:
                arr = arr[0]
            if arr.ndim == 2:
                if arr.shape[0] < HLA_LEN:
                    pad = np.zeros((HLA_LEN - arr.shape[0], ANKH_DIM), np.float32)
                    arr = np.vstack([arr, pad])
                return arr[:HLA_LEN]
    print(f"  [warn] No pre-computed HLA emb for {allele_str}, embedding on-the-fly ...")
    pseq_path = os.path.join(HLA_PSEQ_DIR, f"{tag}.txt")
    if not os.path.exists(pseq_path):
        raise FileNotFoundError(f"No HLA embedding or pseudoseq for {allele_str}.")
    with open(pseq_path) as f:
        pseq = f.read().strip()
    return embed_one_ankh(pseq, HLA_LEN)

def unload_ankh():
    global _ankh_mdl, _ankh_tok
    _ankh_mdl = _ankh_tok = None
    gc.collect()
    if DEVICE.type == "cuda":
        torch.cuda.empty_cache()

class DeepScanCS_PT(nn.Module):
    def __init__(self, feat_dim, win_sizes=WIN_SIZES, num_filter=NUM_FILTER,
                 num_hidden=NUM_HIDDEN, num_class=NUM_CLASS, dropout=0.7):
        super().__init__()
        self.convs = nn.ModuleList([
            nn.Conv1d(feat_dim, num_filter, kernel_size=w, padding=0)
            for w in win_sizes
        ])
        self.pools = nn.ModuleList([nn.AdaptiveMaxPool1d(1) for _ in win_sizes])
        self.classifier = nn.Sequential(
            nn.Dropout(dropout),
            nn.Linear(num_filter * len(win_sizes), num_hidden),
            nn.ReLU(),
            nn.Linear(num_hidden, num_class),
        )
        self.softmax = nn.Softmax(dim=1)

    def forward(self, x):
        branches = [pool(torch.relu(conv(x))).squeeze(-1)
                    for conv, pool in zip(self.convs, self.pools)]
        x = torch.cat(branches, dim=1)
        return self.softmax(self.classifier(x))

def load_pt_model_b(weights_path):
    model = DeepScanCS_PT(feat_dim=ANKH_DIM).to(DEVICE)
    state = torch.load(weights_path, map_location=DEVICE)
    if isinstance(state, dict) and "model_state_dict" in state:
        state = state["model_state_dict"]
    model.load_state_dict(state)
    model.eval()
    return model

MAXSEQ   = 9
CNN_DIM  = 256
ATTN_DIM = 256

class PeptideHLACrossAttention(nn.Module):
    def __init__(self, d_model, num_heads=8, dropout=0.1, attn_dim=ATTN_DIM):
        super().__init__()
        assert attn_dim % num_heads == 0
        self.num_heads = num_heads
        self.d_head    = attn_dim // num_heads
        self.attn_dim  = attn_dim
        self.W_q = nn.Linear(d_model, attn_dim, bias=False)
        self.W_k = nn.Linear(d_model, attn_dim, bias=False)
        self.W_v = nn.Linear(d_model, attn_dim, bias=False)
        self.W_o = nn.Linear(attn_dim, d_model, bias=False)
        self.attn_drop  = nn.Dropout(dropout)
        self.layer_norm = nn.LayerNorm(d_model, eps=1e-6)
        self.scale      = math.sqrt(self.d_head)

    def forward(self, peptide, hla):
        B = peptide.size(0)
        def split_heads(x, seq_len):
            return x.view(B, seq_len, self.num_heads, self.d_head).transpose(1, 2)
        Q = split_heads(self.W_q(peptide), MAXSEQ)
        K = split_heads(self.W_k(hla),     HLA_LEN)
        V = split_heads(self.W_v(hla),     HLA_LEN)
        scores  = torch.matmul(Q, K.transpose(-2, -1)) / self.scale
        weights = self.attn_drop(F.softmax(scores, dim=-1))
        context = torch.matmul(weights, V).transpose(1, 2).contiguous()
        context = context.view(B, MAXSEQ, self.attn_dim)
        out = self.layer_norm(peptide + self.W_o(context))
        return out, weights

class PositionAttention(nn.Module):
    def __init__(self, num_positions=9):
        super().__init__()
        self.score = nn.Linear(num_positions, num_positions, bias=True)

    def forward(self, x):
        w = F.softmax(self.score(x.mean(dim=-1)), dim=-1).unsqueeze(-1)
        return x * w

class DeepScanCrossAttn(nn.Module):
    def __init__(self, maxseq=MAXSEQ, emb_dim=ANKH_DIM,
                 window_sizes=WIN_SIZES, num_filters=NUM_FILTER,
                 num_hidden=NUM_HIDDEN, num_heads=8, attn_dropout=0.1):
        super().__init__()
        self.window_sizes = window_sizes
        self.cross_attn   = PeptideHLACrossAttention(
            d_model=emb_dim, num_heads=num_heads,
            dropout=attn_dropout, attn_dim=ATTN_DIM)
        self.pos_attn     = PositionAttention(num_positions=maxseq)
        self.proj_cnn     = nn.Linear(emb_dim, CNN_DIM)
        self.convs = nn.ModuleList([
            nn.Conv2d(1, num_filters, kernel_size=(1, ws), padding=0)
            for ws in window_sizes])
        self.pools = nn.ModuleList([nn.AdaptiveMaxPool2d((1, 1)) for _ in window_sizes])
        self.dropout = nn.Dropout(0.7)
        self.fc1 = nn.Linear(num_filters * len(window_sizes), num_hidden)
        self.fc2 = nn.Linear(num_hidden, NUM_CLASS)

    def forward(self, peptide, hla):
        x, _ = self.cross_attn(peptide, hla)
        x     = self.pos_attn(x)
        x     = self.proj_cnn(x)
        x     = x.unsqueeze(1)
        outs  = [pool(F.relu(conv(x))).view(x.size(0), -1)
                 for conv, pool in zip(self.convs, self.pools)]
        x = self.dropout(torch.cat(outs, dim=1))
        return self.fc2(F.relu(self.fc1(x)))

def load_pt_model_c(weights_path):
    model = DeepScanCrossAttn().to(DEVICE)
    state = torch.load(weights_path, map_location=DEVICE)
    if isinstance(state, dict) and "model_state_dict" in state:
        state = state["model_state_dict"]
    model.load_state_dict(state)
    model.eval()
    return model

def predict_pt_b(model, X_np, batch_size=256):
    scores = []
    with torch.no_grad():
        for i in range(0, len(X_np), batch_size):
            xb = torch.tensor(X_np[i:i+batch_size], dtype=torch.float32).to(DEVICE)
            xb = xb.permute(0, 2, 1)
            scores.append(model(xb)[:, 1].cpu().numpy())
    return np.concatenate(scores)

def predict_pt_c(model, X_pep, X_hla, batch_size=256):
    scores = []
    with torch.no_grad():
        for i in range(0, len(X_pep), batch_size):
            pb = torch.tensor(X_pep[i:i+batch_size], dtype=torch.float32).to(DEVICE)
            hb = torch.tensor(X_hla[i:i+batch_size], dtype=torch.float32).to(DEVICE)
            logits = model(pb, hb)
            scores.append(F.softmax(logits, dim=-1)[:, 1].cpu().numpy())
    return np.concatenate(scores)

from tensorflow.keras import layers, Model as KModel

class DeepScanCS_TF(KModel):
    def __init__(self, maxseq, num_feat,
                 window_sizes=None, num_filters=None, num_hidden=None):
        super().__init__()
        self.window_sizes = window_sizes if window_sizes is not None else WIN_SIZES
        num_filters = num_filters if num_filters is not None else NUM_FILTER
        num_hidden  = num_hidden  if num_hidden  is not None else NUM_HIDDEN
        self._maxseq = maxseq

        self.conv2d  = []
        self.maxpool = []
        self.flatten = []
        for ws in self.window_sizes:
            self.conv2d.append(layers.Conv2D(
                filters=num_filters, kernel_size=(1, ws),
                activation=tf.nn.relu, padding='valid',
                bias_initializer=tf.constant_initializer(0.1),
                kernel_initializer=tf.keras.initializers.GlorotUniform()))
            self.maxpool.append(layers.MaxPooling2D(
                pool_size=(1, maxseq - ws + 1), strides=(1, maxseq), padding='valid'))
            self.flatten.append(layers.Flatten())
        self.dropout = layers.Dropout(rate=0.7)
        self.fc1 = layers.Dense(
            num_hidden, activation=tf.nn.relu,
            bias_initializer=tf.constant_initializer(0.1),
            kernel_initializer=tf.keras.initializers.GlorotUniform())
        self.fc2 = layers.Dense(
            NUM_CLASS, activation='softmax',
            kernel_regularizer=tf.keras.regularizers.l2(1e-3))

    def call(self, x, training=False):
        _x = []
        for i in range(len(self.window_sizes)):
            x_conv = self.conv2d[i](x)
            x_maxp = self.maxpool[i](x_conv)
            x_flat = self.flatten[i](x_maxp)
            _x.append(x_flat)
        x = layers.Concatenate()(_x)
        x = self.dropout(x, training=training)
        x = self.fc1(x)
        x = self.fc2(x)
        return x

print("\npHLADeep-Pep (stratA)")

def run_strat_a(df):
    if not os.path.isfile(WEIGHTS_A):
        print(f"  [SKIP] weights not found: {WEIGHTS_A}")
        return None
    seqs  = df["peptide"].tolist()
    print(f"  Embedding {len(seqs)} peptides ...")
    X     = embed_peptides_ankh(seqs)
    X_in  = X[:, np.newaxis, :, :]
    model = DeepScanCS_TF(PEP_LEN, ANKH_DIM)
    dummy = tf.zeros((1, 1, PEP_LEN, ANKH_DIM), dtype=tf.float32)
    _     = model(dummy, training=False)
    model.load_weights(WEIGHTS_A)
    batch_size = 256
    preds = []
    for i in range(0, len(X_in), batch_size):
        xb  = tf.constant(X_in[i:i+batch_size], dtype=tf.float32)
        out = model(xb, training=False).numpy()
        preds.append(out)
    pred   = np.concatenate(preds, axis=0)
    scores = pred[:, 1]
    del X, X_in, model, dummy
    gc.collect()
    tf.keras.backend.clear_session()
    return scores

scores_a = run_strat_a(df)
if scores_a is not None:
    print(f"  Score range: {scores_a.min():.4f} – {scores_a.max():.4f}")
    print(f"  Pos mean: {scores_a[y_true==1].mean():.4f}  Neg mean: {scores_a[y_true==0].mean():.4f}")

unload_ankh()
gc.collect()
if DEVICE.type == "cuda":
    torch.cuda.empty_cache()

print("\npHLADeep-Concat (stratB_concat)")

def run_strat_b(df):
    if not os.path.isfile(WEIGHTS_B):
        print(f"  [SKIP] weights not found: {WEIGHTS_B}")
        return None
    seqs    = df["peptide"].tolist()
    alleles = df["allele"].tolist()
    print(f"  Embedding {len(seqs)} peptides ...")
    X_pep = embed_peptides_ankh(seqs)
    hla_cache = {}
    for a in set(alleles):
        hla_cache[a] = load_hla_ankh_emb(a)
    X_hla = np.array([hla_cache[a] for a in alleles], dtype=np.float32)
    X_cat = np.concatenate([X_pep, X_hla], axis=1)
    model  = load_pt_model_b(WEIGHTS_B)
    scores = predict_pt_b(model, X_cat)
    del X_pep, X_hla, X_cat, model
    gc.collect()
    if DEVICE.type == "cuda":
        torch.cuda.empty_cache()
    return scores

scores_b = run_strat_b(df)
if scores_b is not None:
    print(f"  Score range: {scores_b.min():.4f} – {scores_b.max():.4f}")
    print(f"  Pos mean: {scores_b[y_true==1].mean():.4f}  Neg mean: {scores_b[y_true==0].mean():.4f}")

unload_ankh()

print("\npHLADeep-Attn (stratC)")

def run_strat_c(df):
    if not os.path.isfile(WEIGHTS_C):
        print(f"  [SKIP] weights not found: {WEIGHTS_C}")
        return None
    seqs    = df["peptide"].tolist()
    alleles = df["allele"].tolist()
    print(f"  Embedding {len(seqs)} peptides ...")
    X_pep = embed_peptides_ankh(seqs)
    hla_cache = {}
    for a in set(alleles):
        hla_cache[a] = load_hla_ankh_emb(a)
    X_hla = np.array([hla_cache[a] for a in alleles], dtype=np.float32)
    model  = load_pt_model_c(WEIGHTS_C)
    scores = predict_pt_c(model, X_pep, X_hla)
    del X_pep, X_hla, model
    gc.collect()
    if DEVICE.type == "cuda":
        torch.cuda.empty_cache()
    return scores

scores_c = run_strat_c(df)
if scores_c is not None:
    print(f"  Score range: {scores_c.min():.4f} – {scores_c.max():.4f}")
    print(f"  Pos mean: {scores_c[y_true==1].mean():.4f}  Neg mean: {scores_c[y_true==0].mean():.4f}")

unload_ankh()

print("\nNetMHCpan 4.2")

def fmt_netmhc(allele_str):
    return allele_str.replace("*", "")

def _parse_netmhc_text(text, niner_rows, scores, allele, n_sub):
    header   = None
    rank_col = None
    id_col   = None
    pep_col  = None
    parsed   = 0

    for line in text.splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith("#") or stripped.startswith("-"):
            continue
        cols = stripped.split()

        if header is None:
            rank_tokens = {"%Rank_EL", "EL-rank", "EL_rank", "Rank_EL", "%Rank", "EL_%Rank"}
            if any(c in rank_tokens for c in cols):
                header = cols
                for ci, h in enumerate(header):
                    if h in rank_tokens:
                        rank_col = ci
                    if h in ("Identity", "ID", "Peptide_id", "Name"):
                        id_col = ci
                    if h in ("Peptide", "peptide"):
                        pep_col = ci
                continue

        if header is None:
            continue

        try:
            int(cols[0])
        except (ValueError, IndexError):
            continue

        max_col = max(rank_col,
                      id_col  if id_col  is not None else 0,
                      pep_col if pep_col is not None else 0)
        if len(cols) <= max_col:
            continue

        try:
            rank_el = float(cols[rank_col])
            raw_id  = None
            if id_col is not None:
                raw_id = cols[id_col]
            if raw_id is None or not str(raw_id).startswith("row"):
                raw_id = next((c for c in cols if str(c).startswith("row")), None)
            if raw_id is None:
                continue
            ri = int(str(raw_id)[3:])
            scores[ri] = 1.0 - rank_el / 100.0
            parsed += 1
        except Exception:
            continue

    print(f"  {allele}: parsed {parsed}/{n_sub}")
    return parsed, header


def run_netmhcpan(df):
    if not os.path.isfile(NETMHCPAN):
        print(f"  [SKIP] Not found: {NETMHCPAN}")
        return None

    scores = np.full(len(df), np.nan)
    for allele in df["allele"].unique():
        sub         = df[df["allele"] == allele]
        afmt        = fmt_netmhc(allele)
        row_indices = sub.index.tolist()
        niner_rows  = [ri for ri in row_indices if len(df.at[ri, 'peptide']) == 9]
        skipped     = len(row_indices) - len(niner_rows)
        if skipped:
            print(f"  {allele}: skipping {skipped} non-9-mer peptides")
        if not niner_rows:
            continue

        fpath  = None
        xlsout = None
        try:
            with tempfile.NamedTemporaryFile("w", suffix=".fasta", delete=False, dir="/tmp") as f:
                for ri in niner_rows:
                    f.write(f">row{ri}\n{df.at[ri, 'peptide']}\n")
                fpath = f.name

            fd, xlsout = tempfile.mkstemp(suffix=".xls", dir="/tmp")
            os.close(fd)

            cmd = [NETMHCPAN, "-a", afmt, "-f", fpath, "-BA", "-xls", "-xlsfile", xlsout]
            res = subprocess.run(cmd, capture_output=True, text=True, timeout=600)

            if res.returncode != 0:
                print(f"  netMHCpan error for {allele} (rc={res.returncode}): {res.stderr[:400]}")
                continue

            parsed = 0
            if os.path.isfile(xlsout) and os.path.getsize(xlsout) > 0:
                with open(xlsout) as xf:
                    xls_text = xf.read()
                parsed, hdr = _parse_netmhc_text(xls_text, niner_rows, scores, allele, len(sub))
            else:
                parsed, hdr = _parse_netmhc_text(res.stdout, niner_rows, scores, allele, len(sub))

            if parsed == 0:
                all_lines = [l for l in res.stdout.splitlines() if l.strip()]
                print(f"  [DEBUG] stdout ({len(all_lines)} lines):")
                for sl in all_lines:
                    print(f"    {sl}")

        except Exception as e:
            import traceback
            print(f"  Error for {allele}: {e}")
            traceback.print_exc()
        finally:
            for p in [fpath, xlsout]:
                if p and os.path.isfile(p):
                    try:
                        os.unlink(p)
                    except OSError:
                        pass

    valid = ~np.isnan(scores)
    print(f"  Scored {valid.sum()}/{len(df)}")
    return scores

scores_nm = run_netmhcpan(df)

print("\nMHCflurry 2.0")

def run_mhcflurry(df):
    try:
        from mhcflurry import Class1AffinityPredictor
    except ImportError:
        print("  [SKIP] mhcflurry not installed")
        return None

    try:
        predictor = Class1AffinityPredictor.load()
    except Exception as e:
        print(f"  [SKIP] Could not load predictor: {e}")
        return None

    scores = np.full(len(df), np.nan)
    for allele in df["allele"].unique():
        sub      = df[df["allele"] == allele]
        row_idxs = sub.index.tolist()
        peps     = sub["peptide"].tolist()
        valid    = [(ri, p) for ri, p in zip(row_idxs, peps) if 8 <= len(p) <= 11]
        if not valid:
            continue
        vi, vp = zip(*valid)
        vi, vp = list(vi), list(vp)

        try:
            result = predictor.predict(peptides=vp, alleles=[allele] * len(vp))
            if isinstance(result, np.ndarray):
                aff_nm = result.astype(float)
            elif hasattr(result, "values") and not hasattr(result, "columns"):
                aff_nm = result.values.astype(float)
            else:
                col = None
                for cname in ("affinity", "mhcflurry_affinity", "prediction",
                              "mhcflurry_affinity_percentile"):
                    if cname in result.columns:
                        col = cname
                        break
                if col is None:
                    print(f"  {allele}: unknown columns {result.columns.tolist()}")
                    continue
                aff_nm = result[col].values.astype(float)
            sc = 1.0 - np.log(np.clip(aff_nm, 1.0, 50000.0)) / np.log(50000.0)
            sc = np.clip(sc, 0.0, 1.0)
            for j, ri in enumerate(vi):
                scores[ri] = sc[j]
            print(f"  {allele}: scored {len(vi)}/{len(sub)}")
        except Exception as e:
            import traceback
            print(f"  {allele}: error — {e}")
            traceback.print_exc()

    valid_mask = ~np.isnan(scores)
    print(f"  Scored {valid_mask.sum()}/{len(df)} total")
    return scores

scores_mf = run_mhcflurry(df)

print("\nSaving scores")

method_scores = {}
if scores_a  is not None: method_scores["pHLADeep-Pep"]    = scores_a
if scores_b  is not None: method_scores["pHLADeep-Concat"] = scores_b
if scores_c  is not None: method_scores["pHLADeep-Attn"]   = scores_c
if scores_nm is not None and not np.isnan(scores_nm).all():
    method_scores["NetMHCpan 4.2"] = scores_nm
if scores_mf is not None and not np.isnan(scores_mf).all():
    method_scores["MHCflurry 2.0"] = scores_mf

pkl_out = os.path.join(RESULTS_DIR, "casestudy_scores.pkl")
with open(pkl_out, "wb") as f:
    pickle.dump({"df": df, "y_true": y_true, "scores": method_scores}, f)
print(f"  Saved: {pkl_out}")

print("\nPlotting ROC figure")

arial_ok = any("Arial" in f.name for f in fm.fontManager.ttflist)
matplotlib.rcParams.update({
    "font.family":       "Arial" if arial_ok else "sans-serif",
    "font.size":         7,
    "axes.linewidth":    0.75,
    "xtick.major.width": 0.75,
    "ytick.major.width": 0.75,
    "xtick.major.size":  3,
    "ytick.major.size":  3,
    "xtick.direction":   "out",
    "ytick.direction":   "out",
    "pdf.fonttype":      42,
    "ps.fonttype":       42,
})

METHOD_ORDER = ["pHLADeep-Pep", "pHLADeep-Concat", "pHLADeep-Attn",
                "NetMHCpan 4.2", "MHCflurry 2.0"]

FIG_W = 88 / 25.4
FIG_H = 80 / 25.4
fig, ax = plt.subplots(figsize=(FIG_W, FIG_H))
auc_vals    = {}
metric_vals = {}

for mth in METHOD_ORDER:
    if mth not in method_scores:
        continue
    sc   = method_scores[mth].astype(float)
    mask = ~np.isnan(sc)
    n_pos_scored = int(y_true[mask].sum())
    n_neg_scored = int((y_true[mask] == 0).sum())
    if mask.sum() < 2 or len(np.unique(y_true[mask])) < 2:
        continue
    fpr, tpr, thresholds = roc_curve(y_true[mask], sc[mask])
    roc_auc              = sk_auc(fpr, tpr)
    auc_vals[mth]        = roc_auc

    youden_idx = np.argmax(tpr - fpr)
    opt_thr    = thresholds[youden_idx]
    y_pred     = (sc[mask] >= opt_thr).astype(int)
    yt         = y_true[mask].astype(int)
    tn, fp, fn, tp = confusion_matrix(yt, y_pred, labels=[0, 1]).ravel()
    sens = tp / (tp + fn) if (tp + fn) > 0 else 0.0
    spec = tn / (tn + fp) if (tn + fp) > 0 else 0.0
    acc  = (tp + tn) / len(yt)
    mcc  = matthews_corrcoef(yt, y_pred)
    metric_vals[mth] = {"Sens": sens, "Spec": spec, "Acc": acc,
                        "MCC": mcc, "AUC": roc_auc,
                        "N_pos": n_pos_scored, "N_neg": n_neg_scored}

    ax.plot(fpr, tpr, color=COLORS[mth], linestyle=LSTYLES[mth],
            linewidth=1.0, label=f"{mth}  (AUC = {roc_auc:.3f})")

ax.plot([0, 1], [0, 1], color="#AAAAAA", linestyle=":", linewidth=0.6, zorder=0)
ax.set_xlabel("False Positive Rate", fontsize=7)
ax.set_ylabel("True Positive Rate",  fontsize=7)
ax.set_title("Cancer neoepitope–HLA presentation\n(Gurung et al. 2023, Nat. Biotechnol.)",
             fontsize=7, pad=4)
ax.set_xlim(-0.02, 1.02)
ax.set_ylim(-0.02, 1.05)
ax.xaxis.set_major_locator(ticker.MultipleLocator(0.2))
ax.yaxis.set_major_locator(ticker.MultipleLocator(0.2))
ax.tick_params(labelsize=6)
ax.spines["top"].set_visible(False)
ax.spines["right"].set_visible(False)
ax.legend(fontsize=5.5, frameon=False, loc="lower right",
          handlelength=1.5, handletextpad=0.4, labelspacing=0.3)

plt.tight_layout(pad=0.5)
for ext in [".png", ".pdf"]:
    out = os.path.join(RESULTS_DIR, f"fig_casestudy_roc{ext}")
    fig.savefig(out, dpi=300, bbox_inches="tight")
    print(f"  Saved: {out}")
plt.close()

print("\nRESULTS")
print(f"Positives: {int(y_true.sum())} MS-confirmed neoepitope-HLA pairs")
print(f"Negatives: {int((y_true == 0).sum())} random decoys (1:{DECOY_RATIO})")
print(f"Alleles  : {sorted(df[df['label']==1]['allele'].unique())}")
print()
print(f"  {'Method':<22}  {'Sens':>6}  {'Spec':>6}  {'Acc':>6}  {'MCC':>6}  {'AUC':>6}")
print("  " + "-" * 60)
for mth in METHOD_ORDER:
    if mth in metric_vals:
        m = metric_vals[mth]
        print(f"  {mth:<22}  {m['Sens']:.4f}  {m['Spec']:.4f}  "
              f"{m['Acc']:.4f}  {m['MCC']:.4f}  {m['AUC']:.4f}")
    else:
        print(f"  {mth:<22}  — (not scored)")

rows = []
for mth in METHOD_ORDER:
    if mth in metric_vals:
        m = metric_vals[mth]
        rows.append({
            "Method": mth,
            "Sens":   round(m["Sens"], 4),
            "Spec":   round(m["Spec"], 4),
            "Acc":    round(m["Acc"],  4),
            "MCC":    round(m["MCC"],  4),
            "AUC":    round(m["AUC"],  4),
            "N_pos":  m["N_pos"],
            "N_neg":  m["N_neg"],
            "Dataset": "Gurung 2023 (Nat. Biotechnol.)"
        })
pd.DataFrame(rows).to_csv(os.path.join(RESULTS_DIR, "casestudy_results.csv"), index=False)
print(f"\nAll outputs -> {RESULTS_DIR}")
